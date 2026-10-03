#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera a PLANILHA DE ESTUDO viva (.xlsx) — versão com estética da marca.

Abas:
  Conhecimentos Básicos / Específicos  1 linha por item: status, fases (datas + questões), anotações
  Resumo                               cobertura e % de acerto por disciplina, fase e bloco
  Edital                               disciplinas × tipo (Decoreba/Raciocínio) × ciclos em que entraram
  Ciclos                               ciclos de estudo em blocos de 0' a 1h (listas com as disciplinas)
  Controle                             dia a dia de cada ciclo: um 'x' por matéria estudada, cor = volta
  Evolução Semanal                     certas/resolvidas por disciplina, semana a semana até a prova
  Como usar
Compatível com Google Sheets (sem DataBar; referências a outras abas na formatação via INDIRECT).
"""
import argparse, datetime as dt, json, os, re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference

_LOGO_DEFAULT = Path(__file__).resolve().parent.parent / "assets" / "logo.png"
LOGO = os.environ.get("EDITAL_LOGO") or str(_LOGO_DEFAULT)  # opcional: marca própria do usuário

AZUL = "2563EB"; AZUL_DARK = "1D4ED8"; AZUL_LIGHT = "60A5FA"; SLATE = "1E293B"
ZEBRA = "F5F8FD"; BRANCO = "FFFFFF"
F_VERDE = "C6EFCE"; F_AZUL = "DCE9FB"; F_AMARELO = "FFF1C2"; F_VERM = "FBE0E0"
F_INC_ALTA = "F9C9C9"; F_INC_MED = "FDE7B5"; F_INC_BAIXA = "E7EDF5"
F_DECOREBA = "CFE2F3"; F_RACIOCINIO = "FFF2CC"   # mesmas cores da legenda da planilha antiga
# uma cor por volta do ciclo (volta 7 reaproveita a 1ª, e assim por diante)
F_VOLTAS = ["F4A6A6", "9FC5E8", "B6D7A8", "D5C4F0", "FFE599", "D9D9D9"]

STATUS = ["A estudar", "Estudando", "Estudado", "Revisão 1", "Revisão 2", "Revisão 3"]
INCID = ["Sem dado", "Alta", "Média", "Baixa"]
PRIOR = ["Alta", "Média", "Baixa"]
TIPOS = ["Decoreba", "Raciocínio"]
META_ACERTO = 0.70  # abaixo disso a % de acerto fica vermelha; a partir disso, verde
F_ACERTO_OK = "C6EFCE"; F_ACERTO_RUIM = "F9C9C9"

# colunas fixas das abas de itens (nome, largura)
BASE = [("Disciplina", 30), ("Nº", 7), ("Nível", 7), ("Conteúdo", 74),
        ("Status", 13), ("Incidência", 12), ("Prioridade", 11)]
# grupos de fases à direita (rótulo do grupo, tem datas?)
FASES = [("Fase 1", True), ("Fase 2", True), ("Questões Fase 3", False)]

N_CICLOS = 6; BLOCOS_CICLO = 4; SLOTS_BLOCO = 3
SLOTS = BLOCOS_CICLO * SLOTS_BLOCO
DIAS_CONTROLE = 45                 # colunas de dias por ciclo na aba Controle
SEMANAS = 15                       # 05/10/2026 → semana da prova (17/01/2027)
INICIO_SEMANAS = dt.date(2026, 10, 5)

THIN = Side(style="thin", color="DCE3EC")
SEP = Side(style="medium", color=AZUL_LIGHT)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_ROW = 3

BLOCOS = [("basicos", "Conhecimentos Básicos"), ("especificos", "Conhecimentos Específicos")]
ROTULO_BLOCO = {"basicos": "Básicas", "especificos": "Específicas"}
TITLE_SMALL = {"de", "da", "do", "dos", "das", "e", "em", "a", "o", "com", "para", "à", "ao"}


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _nome_curto(nome):
    """'LÍNGUA PORTUGUESA' → 'Língua Portuguesa' (para listas e ciclos)."""
    if not nome.isupper():
        return nome
    ws = nome.lower().split()
    return " ".join(w if i and w in TITLE_SMALL else w[:1].upper() + w[1:] for i, w in enumerate(ws))


def _tipo_sugerido(nome):
    """Sugestão inicial, no critério da planilha antiga: quase tudo Decoreba;
    Raciocínio para o que é interpretação/aplicação prática."""
    u = nome.upper()
    return "Raciocínio" if re.search(r"INGLESA|DIVULGA|RACIOC|L[ÓO]GICA|MATEM|ESTAT", u) else "Decoreba"


# ---------------------------------------------------------------- layout das colunas
def _layout():
    """Mapa de colunas das abas de itens: fases[i] = {nome, ini, fim, certas, resol, pct}."""
    col = len(BASE) + 1
    fases = []
    for nome, datas in FASES:
        f = {"nome": nome, "primeira": col}
        if datas:
            f["ini"], f["fim"] = col, col + 1; col += 2
        f["certas"], f["resol"], f["pct"] = col, col + 1, col + 2; col += 3
        f["ultima"] = col - 1
        fases.append(f)
    return fases, col, col + 1     # fases, coluna Anotações, coluna oculta _feito


FASES_COLS, COL_ANOT, COL_FEITO = _layout()


def _abas(data):
    """[(nome da aba, chave do bloco, disciplinas)] — uma aba por bloco; sem blocos, aba única."""
    discs = data.get("disciplinas", [])
    if not any(d.get("bloco") for d in discs):
        return [("Verticalizado", "", discs)]
    return [(nome, key, [d for d in discs if d.get("bloco") == key])
            for key, nome in BLOCOS if any(d.get("bloco") == key for d in discs)]


def build(data, out):
    wb = Workbook(); wb.remove(wb.active)
    abas = []
    for nome, key, discs in _abas(data):
        ws = wb.create_sheet(nome)
        abas.append((nome, key, discs, _aba_itens(ws, data, discs)))
    _resumo(wb, data, abas)
    lista = _edital(wb, abas)
    _ciclos(wb, lista)
    _controle(wb)
    _evolucao(wb, abas)
    _ajuda(wb)
    wb.save(out)
    return abas


# ---------------------------------------------------------------- abas de itens
def _banner(ws, data, bloco, ultima_col):
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=ultima_col)
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=len(BASE))
    for col in range(2, ultima_col + 1):
        ws.cell(1, col).fill = _fill(AZUL)
        ws.cell(2, col).fill = _fill(AZUL_DARK)
    ws["B1"] = f"Edital Verticalizado — {bloco}" if bloco and bloco != "Verticalizado" \
        else "Edital Verticalizado — plano de estudo"
    ws["B1"].font = Font(bold=True, color=BRANCO, size=15)
    ws["B1"].alignment = Alignment(vertical="center", indent=1)
    sub = " · ".join(x for x in [data.get("concurso"), data.get("cargo")] if x)
    ws["B2"] = sub or "Plano de estudo"
    ws["B2"].font = Font(color="DBE7FF", size=10.5)
    ws["B2"].alignment = Alignment(vertical="center", indent=1)
    for f in FASES_COLS:   # rótulo de cada grupo de fase sobre suas colunas
        ws.merge_cells(start_row=2, start_column=f["primeira"], end_row=2, end_column=f["ultima"])
        c = ws.cell(2, f["primeira"], f["nome"])
        c.font = Font(bold=True, color=BRANCO, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
    if not LOGO:
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(LOGO); img.height = 44; img.width = 44
        ws.add_image(img, "A1")
    except Exception:
        pass  # sem logo: o título do banner se sustenta sozinho


def _header(ws):
    nomes = [n for n, _ in BASE]
    for f in FASES_COLS:
        nomes += (["Início", "Conclusão"] if "ini" in f else []) + ["Certas", "Resolvidas", "Porcentagem"]
    nomes.append("Anotações")
    for c, name in enumerate(nomes, start=1):
        cell = ws.cell(HEAD_ROW, c, name)
        cell.fill = _fill(AZUL); cell.font = Font(bold=True, color=BRANCO, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX
    ws.cell(HEAD_ROW, COL_FEITO, "_feito")


def _aba_itens(ws, data, discs):
    ws.sheet_view.showGridLines = False; ws.sheet_view.zoomScale = 110
    ws.sheet_properties.tabColor = AZUL
    _banner(ws, data, ws.title, COL_ANOT); _header(ws)

    centro = set(range(2, COL_ANOT)) - {4}  # tudo centralizado, menos Disciplina/Conteúdo/Anotações
    r = HEAD_ROW + 1
    for di, disc in enumerate(discs):
        nome = disc.get("nome", "")
        block = _fill(ZEBRA if di % 2 == 0 else BRANCO)
        first = True
        for it in disc.get("itens", []):
            nivel = int(it.get("nivel", 1))
            vals = {1: nome, 2: it.get("numero", ""), 3: nivel,
                    4: ("   " * (nivel - 1)) + it.get("texto", ""), 5: "A estudar", 6: "Sem dado"}
            for f in FASES_COLS:
                c, s = L(f["certas"]), L(f["resol"])
                vals[f["pct"]] = f'=IF(N({s}{r})=0,"",{c}{r}/{s}{r})'
            for c in range(1, COL_ANOT + 1):
                cell = ws.cell(r, c, vals.get(c))
                cell.fill = block
                cell.border = Border(left=THIN, right=THIN, top=SEP if first else THIN, bottom=THIN)
                if c in centro:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(vertical="center", wrap_text=(c in (4, COL_ANOT)))
            for f in FASES_COLS:
                ws.cell(r, f["pct"]).number_format = "0.00%"
                if "ini" in f:
                    ws.cell(r, f["ini"]).number_format = "dd/mm/yyyy"
                    ws.cell(r, f["fim"]).number_format = "dd/mm/yyyy"
            # assunto (nível 1) em negrito
            if nivel == 1:
                ws.cell(r, 1).font = Font(bold=True, color=SLATE, size=10)
                ws.cell(r, 4).font = Font(bold=True, color=SLATE)
            else:
                ws.cell(r, 1).font = Font(color="64748B", size=9)
                ws.cell(r, 4).font = Font(color="334155")
            ws.cell(r, COL_FEITO, f'=IF(OR($E{r}="Estudado",LEFT($E{r},7)="Revisão"),1,0)').fill = block
            first = False
            r += 1
    last = r - 1

    for c, (_, w) in enumerate(BASE, start=1):
        ws.column_dimensions[L(c)].width = w
    for f in FASES_COLS:
        if "ini" in f:
            ws.column_dimensions[L(f["ini"])].width = 11
            ws.column_dimensions[L(f["fim"])].width = 11
        ws.column_dimensions[L(f["certas"])].width = 8
        ws.column_dimensions[L(f["resol"])].width = 11
        ws.column_dimensions[L(f["pct"])].width = 13
    ws.column_dimensions[L(COL_ANOT)].width = 26
    ws.column_dimensions[L(COL_FEITO)].width = 6
    ws.column_dimensions[L(COL_FEITO)].hidden = True
    ws.freeze_panes = "E4"
    ws.auto_filter.ref = f"A{HEAD_ROW}:{L(COL_ANOT)}{last}"

    h0 = HEAD_ROW + 1

    def dv(col, opts):
        d = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        d.add(f"{col}{h0}:{col}{last}"); ws.add_data_validation(d)
    dv("E", STATUS); dv("F", INCID); dv("G", PRIOR)

    for f in FASES_COLS:
        c, s = L(f["certas"]), L(f["resol"])
        _valida_questoes(ws, c, s, h0, last)
        _cor_acerto(ws, f"{L(f['pct'])}{h0}:{L(f['pct'])}{last}", f"{L(f['pct'])}{h0}")
        if "ini" in f:
            i, e = L(f["ini"]), L(f["fim"])
            _valida(ws, f"{i}{h0}:{i}{last}", f"ISNUMBER({i}{h0})",
                    "Informe uma data (dd/mm/aaaa).")
            _valida(ws, f"{e}{h0}:{e}{last}", f'AND(ISNUMBER({e}{h0}),OR({i}{h0}="",{e}{h0}>={i}{h0}))',
                    "Informe uma data (dd/mm/aaaa) igual ou posterior ao início.")

    e = f"E{h0}:E{last}"; f = f"F{h0}:F{last}"
    rules = [
        (e, f'$E{h0}="Estudado"', F_VERDE),
        (e, f'LEFT($E{h0},7)="Revisão"', F_AZUL),
        (e, f'$E{h0}="Estudando"', F_AMARELO),
        (e, f'$E{h0}="A estudar"', F_VERM),
        (f, f'$F{h0}="Alta"', F_INC_ALTA),
        (f, f'$F{h0}="Média"', F_INC_MED),
        (f, f'$F{h0}="Baixa"', F_INC_BAIXA),
    ]
    for rng, formula, color in rules:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], fill=_fill(color)))
    return last


def _valida(ws, rng, formula, msg):
    d = DataValidation(type="custom", formula1=formula, allow_blank=True,
                       showErrorMessage=True, errorTitle="Valor inválido", error=msg)
    d.add(rng); ws.add_data_validation(d)


def _valida_questoes(ws, c, s, first, last):
    """Certas/Resolvidas: inteiros >= 0; Certas não pode passar de Resolvidas."""
    _valida(ws, f"{c}{first}:{c}{last}",
            f'AND(INT({c}{first})={c}{first},{c}{first}>=0,OR({s}{first}="",{c}{first}<={s}{first}))',
            "Informe um número inteiro, no máximo igual ao de questões resolvidas.")
    _valida(ws, f"{s}{first}:{s}{last}",
            f'AND(INT({s}{first})={s}{first},{s}{first}>=0,OR({c}{first}="",{s}{first}>={c}{first}))',
            "Informe um número inteiro, no mínimo igual ao de questões certas.")


def _cor_acerto(ws, rng, cell):
    ok = f"AND(ISNUMBER({cell}),{cell}>={META_ACERTO})"
    ruim = f"AND(ISNUMBER({cell}),{cell}<{META_ACERTO})"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[ok], fill=_fill(F_ACERTO_OK),
                                                   font=Font(bold=True, color="166534")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[ruim], fill=_fill(F_ACERTO_RUIM),
                                                   font=Font(bold=True, color="991B1B")))


def _cabecalho(ws, row, col, nome, larg=None):
    cell = ws.cell(row, col, nome); cell.fill = _fill(AZUL)
    cell.font = Font(bold=True, color=BRANCO); cell.border = BOX
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if larg:
        ws.column_dimensions[L(col)].width = larg
    return cell


def _grupo(ws, row, c1, c2, nome):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    cell = ws.cell(row, c1, nome); cell.fill = _fill(AZUL_DARK)
    cell.font = Font(bold=True, color=BRANCO); cell.alignment = Alignment(horizontal="center")
    for c in range(c1 + 1, c2 + 1):
        ws.cell(row, c).fill = _fill(AZUL_DARK)


def _pct(ws, row, col, certas, resol):
    c = ws.cell(row, col, f'=IF(N({resol})=0,"",{certas}/{resol})')
    c.number_format = "0.00%"; c.alignment = Alignment(horizontal="center")
    return c


def _titulo(ws, titulo, nota, cor):
    ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = cor
    ws["A1"] = titulo; ws["A1"].font = Font(bold=True, size=15, color=SLATE)
    if nota:
        ws["A2"] = nota; ws["A2"].font = Font(italic=True, color="64748B")


# ---------------------------------------------------------------- Resumo
def _resumo(wb, data, abas):
    ws = wb.create_sheet("Resumo")
    _titulo(ws, "Resumo de cobertura",
            " · ".join(x for x in [data.get("concurso"), data.get("cargo")] if x), AZUL_LIGHT)
    grp, hdr_row = 3, 4
    for c, (name, w) in enumerate([("Disciplina", 34), ("Itens", 8), ("Concluídos", 11),
                                   ("% Cobertura", 12), ("Progresso", 24)], start=1):
        _cabecalho(ws, hdr_row, c, name, w)
    # por fase: Certas | Resolvidas | %   e, no fim, o total das fases
    cols_fase = []
    col = 6
    for f in FASES_COLS + [{"nome": "Total (todas as fases)"}]:
        _grupo(ws, grp, col, col + 2, f["nome"])
        for k, (name, w) in enumerate([("Certas", 8), ("Resolvidas", 11), ("% Acerto", 11)]):
            _cabecalho(ws, hdr_row, col + k, name, w)
        cols_fase.append((f, col)); col += 3
    ultima = col - 1

    row = hdr_row + 1
    first_data = row
    faixas = []  # (nome da aba, 1ª linha, última linha) no resumo
    for aba, _, discs, last in abas:
        ref = lambda c: f"'{aba}'!${c}${HEAD_ROW+1}:${c}${last}"
        ini = row
        for disc in discs:
            ws.cell(row, 1, disc.get("nome", ""))
            ws.cell(row, 2, f"=COUNTIF({ref('A')},A{row})")
            ws.cell(row, 3, f"=SUMIFS({ref(L(COL_FEITO))},{ref('A')},A{row})")
            p = ws.cell(row, 4, f"=IF(B{row}=0,0,C{row}/B{row})"); p.number_format = "0%"
            _barra(ws, row)
            soma_c, soma_r = [], []
            for f, c0 in cols_fase:
                if "certas" in f:
                    ws.cell(row, c0, f"=SUMIFS({ref(L(f['certas']))},{ref('A')},$A{row})")
                    ws.cell(row, c0 + 1, f"=SUMIFS({ref(L(f['resol']))},{ref('A')},$A{row})")
                    soma_c.append(f"{L(c0)}{row}"); soma_r.append(f"{L(c0+1)}{row}")
                else:   # total das fases
                    ws.cell(row, c0, "=" + "+".join(soma_c))
                    ws.cell(row, c0 + 1, "=" + "+".join(soma_r))
                _pct(ws, row, c0 + 2, f"{L(c0)}{row}", f"{L(c0+1)}{row}")
            for c in range(1, ultima + 1):
                ws.cell(row, c).border = Border(bottom=THIN)
            row += 1
        faixas.append((aba, ini, row - 1))
    last_data = row - 1

    def _soma(row, rotulo, linhas):
        ws.cell(row, 1, rotulo).font = Font(bold=True)
        for c in [2, 3] + [c0 + k for _, c0 in cols_fase for k in (0, 1)]:
            ws.cell(row, c, "=" + "+".join(f"SUM({L(c)}{a}:{L(c)}{b})" for a, b in linhas)).font = Font(bold=True)
        tp = ws.cell(row, 4, f"=IF(B{row}=0,0,C{row}/B{row})"); tp.number_format = "0%"; tp.font = Font(bold=True)
        _barra(ws, row)
        for _, c0 in cols_fase:
            _pct(ws, row, c0 + 2, f"{L(c0)}{row}", f"{L(c0+1)}{row}").font = Font(bold=True)

    if len(faixas) > 1:
        for aba, a, b in faixas:
            _soma(row, f"Subtotal – {aba}", [(a, b)]); row += 1
    _soma(row, "TOTAL", [(first_data, last_data)])
    for _, c0 in cols_fase:
        _cor_acerto(ws, f"{L(c0+2)}{first_data}:{L(c0+2)}{row}", f"{L(c0+2)}{first_data}")
    ws.freeze_panes = "B5"

    chart = BarChart(); chart.type = "bar"; chart.title = "% de cobertura por disciplina"
    chart.height = max(7, 0.5 * (last_data - first_data + 1)); chart.width = 18
    chart.legend = None
    vals = Reference(ws, min_col=4, min_row=hdr_row, max_row=last_data)
    cats = Reference(ws, min_col=1, min_row=first_data, max_row=last_data)
    chart.add_data(vals, titles_from_data=True); chart.set_categories(cats)
    chart.y_axis.numFmt = "0%"; chart.x_axis.delete = False; chart.y_axis.delete = False
    ws.add_chart(chart, f"A{row + 3}")


def _barra(ws, row):
    # barra em texto: DataBar não sobrevive à importação no Google Sheets
    c = ws.cell(row, 5, f'=REPT("█",ROUND(D{row}*20,0))&REPT("░",20-ROUND(D{row}*20,0))')
    c.font = Font(color=AZUL, name="Consolas")


# ---------------------------------------------------------------- Edital + Ciclos + Controle
def _lista_disciplinas(abas):
    """[(rótulo do bloco, nome curto)] na ordem do edital."""
    return [(ROTULO_BLOCO.get(key, ""), _nome_curto(d.get("nome", "")))
            for _, key, discs, _ in abas for d in discs]


def _ciclo_linha(k):
    """1ª linha do ciclo k (0-based) na aba Ciclos."""
    return 4 + k * (SLOTS_BLOCO + 3)


def _ciclo_slots(k):
    """Células das disciplinas do ciclo k, na ordem de estudo: linha a linha, bloco a bloco."""
    base = _ciclo_linha(k)
    return [f"{L(1 + 2 * b)}{base + 2 + s}" for s in range(SLOTS_BLOCO) for b in range(BLOCOS_CICLO)]


def _ciclo_area(k):
    base = _ciclo_linha(k)
    return f"$A${base + 2}:${L(2 * BLOCOS_CICLO)}${base + 1 + SLOTS_BLOCO}"


def _edital(wb, abas):
    ws = wb.create_sheet("Edital")
    _titulo(ws, "Edital — disciplinas e ciclos",
            "Tipo: sugestão inicial, mude pela lista. As colunas de ciclo marcam 'x' "
            "sozinhas quando a disciplina entra no ciclo (aba Ciclos).", "0EA5E9")
    hdr = 3
    for c, (name, w) in enumerate([("Bloco", 12), ("Disciplina", 46), ("Tipo", 13)], start=1):
        _cabecalho(ws, hdr, c, name, w)
    for k in range(N_CICLOS):
        _cabecalho(ws, hdr, 4 + k, f"Ciclo {k+1:02d}", 10)
    r = hdr + 1
    bloco_ant = None
    for bloco, nome in _lista_disciplinas(abas):
        ws.cell(r, 1, bloco if bloco != bloco_ant else "").font = Font(bold=True, color=SLATE)
        bloco_ant = bloco
        ws.cell(r, 2, nome)
        ws.cell(r, 3, _tipo_sugerido(nome)).alignment = Alignment(horizontal="center")
        for k in range(N_CICLOS):
            c = ws.cell(r, 4 + k, f'=IF(COUNTIF(Ciclos!{_ciclo_area(k)},$B{r})>0,"x","")')
            c.alignment = Alignment(horizontal="center"); c.font = Font(bold=True, color=AZUL)
        for c in range(1, 4 + N_CICLOS):
            ws.cell(r, c).border = Border(bottom=THIN)
        r += 1
    last = r - 1
    d = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=True)
    d.add(f"C{hdr+1}:C{last}"); ws.add_data_validation(d)
    for tipo, cor in (("Decoreba", F_DECOREBA), ("Raciocínio", F_RACIOCINIO)):
        ws.conditional_formatting.add(f"B{hdr+1}:C{last}",
            FormulaRule(formula=[f'$C{hdr+1}="{tipo}"'], fill=_fill(cor)))
    ws.freeze_panes = "C4"
    return f"Edital!$B${hdr+1}:$B${last}"


def _ciclos(wb, lista_disc):
    ws = wb.create_sheet("Ciclos")
    _titulo(ws, "Legenda", "", "0EA5E9")
    ws["A1"].font = Font(bold=True)
    for c, (tipo, cor) in enumerate([("Decoreba", F_DECOREBA), ("Raciocínio", F_RACIOCINIO)], start=2):
        cell = ws.cell(1, c, tipo); cell.fill = _fill(cor); cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
    ws["A2"] = ("Escolha as disciplinas de cada bloco na lista, na ordem em que vai estudar "
                "(linha a linha). A cor vem do Tipo na aba Edital.")
    ws["A2"].font = Font(italic=True, color="64748B")
    for c in range(1, 2 * BLOCOS_CICLO + 1):
        ws.column_dimensions[L(c)].width = 19
    lista = DataValidation(type="list", formula1=lista_disc, allow_blank=True,
                           showErrorMessage=True, errorTitle="Disciplina",
                           error="Escolha uma disciplina da lista (aba Edital).")
    ws.add_data_validation(lista)
    edital_b, edital_c = 'INDIRECT("Edital!$B:$B")', 'INDIRECT("Edital!$C:$C")'
    for k in range(N_CICLOS):
        base = _ciclo_linha(k)
        ws.cell(base, 1, f"Ciclo {k+1:02d}:").font = Font(bold=True, size=12, color=SLATE)
        d = ws.cell(base, 2); d.number_format = "dd/mm/yyyy"; d.border = BOX
        ws.cell(base, 3, "← data de início").font = Font(italic=True, color="94A3B8", size=9)
        for b in range(BLOCOS_CICLO):
            c1 = 1 + 2 * b
            ws.merge_cells(start_row=base + 1, start_column=c1, end_row=base + 1, end_column=c1 + 1)
            h = ws.cell(base + 1, c1, "0' a 1h"); h.fill = _fill(AZUL); h.font = Font(bold=True, color=BRANCO)
            h.alignment = Alignment(horizontal="center")
            for s in range(SLOTS_BLOCO):
                r = base + 2 + s
                ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + 1)
                cell = ws.cell(r, c1)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                for c in (c1, c1 + 1):
                    ws.cell(r, c).border = BOX
                lista.add(cell.coordinate)
        area = f"A{base+2}:{L(2*BLOCOS_CICLO)}{base+1+SLOTS_BLOCO}"
        a0 = f"A{base+2}"
        for tipo, cor in (("Decoreba", F_DECOREBA), ("Raciocínio", F_RACIOCINIO)):
            ws.conditional_formatting.add(area, FormulaRule(
                formula=[f'AND({a0}<>"",COUNTIFS({edital_b},{a0},{edital_c},"{tipo}")>0)'], fill=_fill(cor)))
        for s in range(SLOTS_BLOCO):
            ws.row_dimensions[base + 2 + s].height = 30
    ws.freeze_panes = "A3"


def _controle(wb):
    """Uma grade por ciclo: matérias do ciclo × dias a partir da data de início.
    A pessoa marca 'x' no dia em que estudou; o n-ésimo 'x' da linha ganha a cor da volta n.
    'Volta atual' e 'Próxima matéria' seguem a ordem do ciclo (não se pula matéria)."""
    ws = wb.create_sheet("Controle")
    _titulo(ws, "Controle dos ciclos",
            "Marque 'x' no dia em que estudou a matéria. A cor mostra a volta do ciclo; "
            "quando todas fecham a volta, a próxima começa com outra cor.", "0EA5E9")
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 8
    d0 = 3                                         # 1ª coluna de dias
    dl = d0 + DIAS_CONTROLE - 1                    # última coluna de dias
    aux = dl + 1                                   # coluna auxiliar (oculta) da ordem
    for c in range(d0, dl + 1):
        ws.column_dimensions[L(c)].width = 6.5
    ws.column_dimensions[L(aux)].hidden = True
    # legenda das voltas
    ws.cell(3, 1, "Voltas:").font = Font(bold=True)
    for j, cor in enumerate(F_VOLTAS):
        cell = ws.cell(3, d0 + j, j + 1); cell.fill = _fill(cor); cell.border = BOX
        cell.alignment = Alignment(horizontal="center"); cell.font = Font(bold=True)

    altura = SLOTS + 5
    for k in range(N_CICLOS):
        top = 5 + k * altura
        base = _ciclo_linha(k)
        ws.merge_cells(start_row=top, start_column=1, end_row=top, end_column=dl)
        t = ws.cell(top, 1, f'="Ciclo {k+1:02d}"&IF(Ciclos!$B${base}="","  (defina a data de início na aba Ciclos)",'
                            f'" — início em "&TEXT(Ciclos!$B${base},"dd/mm/yyyy"))')
        t.fill = _fill(AZUL_DARK); t.font = Font(bold=True, color=BRANCO, size=12)
        _cabecalho(ws, top + 1, 1, "Matéria (ordem do ciclo)"); _cabecalho(ws, top + 1, 2, "Vezes")
        for i, c in enumerate(range(d0, dl + 1)):
            h = _cabecalho(ws, top + 1, c, f'=IF(Ciclos!$B${base}="","",Ciclos!$B${base}+{i})')
            h.number_format = "dd/mm"; h.font = Font(bold=True, color=BRANCO, size=9)
        s0, s1 = top + 2, top + 1 + SLOTS
        for s, origem in enumerate(_ciclo_slots(k)):
            r = s0 + s
            ws.cell(r, 1, f'=IF(Ciclos!{origem}="","",Ciclos!{origem})')
            v = ws.cell(r, 2, f'=IF(A{r}="","",COUNTIF({L(d0)}{r}:{L(dl)}{r},"x"))')
            v.alignment = Alignment(horizontal="center")
            # ordem: vezes*100 + posição no ciclo → o menor é a próxima matéria
            ws.cell(r, aux, f'=IF(A{r}="","",B{r}*100+{s+1})')
            for c in range(1, dl + 1):
                ws.cell(r, c).border = BOX
                if c >= d0:
                    ws.cell(r, c).alignment = Alignment(horizontal="center")
        xs = f"{L(d0)}{s0}:{L(dl)}{s1}"
        for j, cor in enumerate(F_VOLTAS):
            ws.conditional_formatting.add(xs, FormulaRule(formula=[
                f'AND(LOWER({L(d0)}{s0})="x",MOD(COUNTIF(${L(d0)}{s0}:{L(d0)}{s0},"x")-1,{len(F_VOLTAS)})={j})'],
                fill=_fill(cor)))
        ordem = f"{L(aux)}{s0}:{L(aux)}{s1}"
        nomes = f"A{s0}:A{s1}"
        ra, rp = s1 + 1, s1 + 2
        ws.cell(ra, 1, "Volta atual").font = Font(bold=True)
        ws.cell(ra, 2, f'=IF(COUNT({ordem})=0,"",INT(MIN({ordem})/100)+1)').font = Font(bold=True, size=12)
        ws.cell(ra, 2).alignment = Alignment(horizontal="center")
        ws.cell(rp, 1, "Próxima matéria").font = Font(bold=True)
        ws.merge_cells(start_row=rp, start_column=2, end_row=rp, end_column=d0 + 6)
        p = ws.cell(rp, 2, f'=IF(COUNT({ordem})=0,"",INDEX({nomes},MATCH(MIN({ordem}),{ordem},0)))')
        p.font = Font(bold=True, color=AZUL_DARK, size=12)
    ws.freeze_panes = "C1"


# ---------------------------------------------------------------- Evolução semanal
def _evolucao(wb, abas):
    ws = wb.create_sheet("Evolução Semanal")
    _titulo(ws, "Evolução semanal de questões",
            "Ao fim de cada semana, lance as questões que fez nela. "
            "Mude a data da 1ª semana se quiser; as outras seguem de 7 em 7 dias.", AZUL_LIGHT)
    grp, hdr = 4, 5
    _cabecalho(ws, hdr, 1, "Bloco", 12); _cabecalho(ws, hdr, 2, "Disciplina", 40)
    for c in (1, 2):
        ws.cell(grp, c).fill = _fill(AZUL_DARK)
    first = hdr + 1
    r = first
    faixas, bloco_ant = {}, None
    for bloco, nome in _lista_disciplinas(abas):
        ws.cell(r, 1, bloco if bloco != bloco_ant else "").font = Font(bold=True, color=SLATE)
        bloco_ant = bloco
        ws.cell(r, 2, nome)
        faixas.setdefault(bloco, [r, r])[1] = r
        r += 1
    last = r - 1
    totais = []
    if len(faixas) > 1:
        for bloco, (a, b) in faixas.items():
            totais.append((r, f"Subtotal – {bloco}", [(a, b)])); r += 1
    totais.append((r, "TOTAL", [(first, last)]))
    for rr, rot, _ in totais:
        ws.cell(rr, 2, rot).font = Font(bold=True)

    for w in range(SEMANAS):
        c0 = 3 + 3 * w
        ws.merge_cells(start_row=grp, start_column=c0, end_row=grp, end_column=c0 + 2)
        dc = ws.cell(grp, c0, INICIO_SEMANAS if w == 0 else f"={L(c0-3)}{grp}+7")
        dc.number_format = '"Semana de "dd/mm/yyyy'
        dc.fill = _fill(AZUL_DARK if w else AZUL_LIGHT)
        dc.font = Font(bold=True, color=BRANCO); dc.alignment = Alignment(horizontal="center")
        for k, (name, wd) in enumerate([("Certas", 8), ("Resolvidas", 11), ("Porcentagem", 12)]):
            _cabecalho(ws, hdr, c0 + k, name, wd)
        ct, rs, pc = L(c0), L(c0 + 1), L(c0 + 2)
        for rr in range(first, last + 1):
            _pct(ws, rr, c0 + 2, f"{ct}{rr}", f"{rs}{rr}")
            for c in range(c0, c0 + 3):
                ws.cell(rr, c).border = BOX
                ws.cell(rr, c).alignment = Alignment(horizontal="center")
        for rr, _, faixa in totais:
            for c in (ct, rs):
                ws[f"{c}{rr}"] = "=" + "+".join(f"SUM({c}{a}:{c}{b})" for a, b in faixa)
                ws[f"{c}{rr}"].font = Font(bold=True); ws[f"{c}{rr}"].alignment = Alignment(horizontal="center")
            _pct(ws, rr, c0 + 2, f"{ct}{rr}", f"{rs}{rr}").font = Font(bold=True)
        _valida_questoes(ws, ct, rs, first, last)
        _cor_acerto(ws, f"{pc}{first}:{pc}{totais[-1][0]}", f"{pc}{first}")
    ws.freeze_panes = ws.cell(hdr + 1, 3).coordinate


# ---------------------------------------------------------------- Como usar
def _ajuda(wb):
    ws = wb.create_sheet("Como usar")
    ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = "94A3B8"
    linhas = [
        ("Como usar esta planilha", True),
        ("", False),
        ("1. Abas 'Conhecimentos Básicos' e 'Conhecimentos Específicos': cada linha é um item do edital,", False),
        ("   na ordem original. Os básicos são comuns a todos os cargos; os específicos são do seu cargo.", False),
        ("2. 'Status': marque seu progresso (A estudar → Estudado → Revisão 1/2/3).", False),
        ("   As cores e a aba 'Resumo' (% de cobertura) atualizam sozinhas.", False),
        ("3. 'Fase 1' e 'Fase 2': a 1ª passada e a revisão de cada item. Anote a data de início e de", False),
        ("   conclusão e as questões feitas (Certas / Resolvidas). 'Questões Fase 3' é só de questões.", False),
        ("   A 'Porcentagem' calcula sozinha: vermelho abaixo de 70%, verde a partir de 70%.", False),
        ("4. 'Incidência': começa como 'Sem dado' DE PROPÓSITO. Para preencher de verdade, junte os", False),
        ("   assuntos das últimas provas da banca e classifique cada item em Alta/Média/Baixa.", False),
        ("5. 'Resumo': cobertura e % de acerto por disciplina, por fase e por bloco (básicos/específicos).", False),
        ("6. 'Edital': classifique cada disciplina como Decoreba ou Raciocínio. As colunas de ciclo", False),
        ("   marcam sozinhas em quais ciclos cada disciplina entrou.", False),
        ("7. 'Ciclos': para cada ciclo, ponha a data de início e escolha as disciplinas nas listas,", False),
        ("   na ordem em que vai estudar (linha a linha). Estude sempre a próxima, sem pular matéria.", False),
        ("8. 'Controle': marque 'x' no dia em que estudou cada matéria do ciclo. O 1º 'x' da linha fica", False),
        ("   com a cor da volta 1, o 2º com a da volta 2, e assim por diante. 'Volta atual' e", False),
        ("   'Próxima matéria' dizem onde você está no ciclo.", False),
        ("9. 'Evolução Semanal': ao fim de cada semana, lance as questões feitas nela, por disciplina.", False),
        ("", False),
        ("No Google Sheets (computador e celular)", True),
        ("", False),
        ("• Envie este arquivo ao Google Drive e abra com 'Planilhas Google' (Arquivo → Salvar como Planilhas Google).", False),
        ("• O progresso é salvo sozinho na sua conta e aparece igual em todos os aparelhos.", False),
        ("• No celular, instale o app Planilhas Google; toque na célula de Status e escolha na lista.", False),
        ("• Para ver só o que falta: Dados → Criar filtro, e filtre a coluna Status.", False),
    ]
    for i, (txt, bold) in enumerate(linhas, start=1):
        c = ws.cell(i, 1, txt)
        c.font = Font(bold=bold, size=15 if bold else 11, color=SLATE if bold else "334155")
    ws.column_dimensions["A"].width = 100


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json"); ap.add_argument("--out", "-o", required=True)
    ap.add_argument("--sem-logo", action="store_true", help="não inserir logo no banner")
    args = ap.parse_args()
    if args.sem_logo:
        global LOGO
        LOGO = ""
    data = json.loads(Path(args.json).read_text(encoding="utf-8"))
    abas = build(data, args.out)
    print(f"OK -> {args.out}")
    for nome, _, discs, last in abas:
        print(f"  {nome}: {len(discs)} disciplinas, {last - HEAD_ROW} itens")


if __name__ == "__main__":
    main()
