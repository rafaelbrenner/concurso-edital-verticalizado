#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o edital verticalizado em DOCX imprimível/anotável (caixas ☐ por item).

Quando o JSON traz `bloco`, o documento se divide em Conhecimentos Básicos e
Conhecimentos Específicos. Cada disciplina vira uma tabela com espaço para anotar
à mão as questões de cada fase (certas/resolvidas), como na planilha.

Uso: python3 gerar_docx.py edital.json --out edital.docx [--sem-logo]
"""
import argparse, json, os
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_LOGO_DEFAULT = Path(__file__).resolve().parent.parent / "assets" / "logo.png"
LOGO = os.environ.get("EDITAL_LOGO") or str(_LOGO_DEFAULT)  # opcional: marca própria do usuário
AZUL = RGBColor(0x25, 0x63, 0xEB)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
MUT = RGBColor(0x64, 0x74, 0x8B)
BLOCOS = {"basicos": "Conhecimentos Básicos", "especificos": "Conhecimentos Específicos"}
FASES = ["Fase 1", "Fase 2", "Fase 3"]
LARG = [Cm(0.7), Cm(1.1), Cm(9.6), Cm(1.8), Cm(1.8), Cm(1.8)]   # ☐ | Nº | Conteúdo | F1 | F2 | F3


def _sombra(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def _texto(cell, txt, bold=False, size=10, cor=None, centro=False):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    if centro:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); r.bold = bold; r.font.size = Pt(size)
    if cor is not None:
        r.font.color.rgb = cor
    return r


def _repete_cabecalho(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def _tabela_itens(doc, itens):
    t = doc.add_table(rows=1, cols=len(LARG)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, (cell, nome) in enumerate(zip(t.rows[0].cells, ["", "Nº", "Conteúdo"] + FASES)):
        _texto(cell, nome, bold=True, size=9, cor=RGBColor(0xFF, 0xFF, 0xFF), centro=c != 2)
        _sombra(cell, "2563EB")
    _repete_cabecalho(t.rows[0])
    for it in itens:
        nivel = int(it.get("nivel", 1))
        cells = t.add_row().cells
        _texto(cells[0], "☐", size=11, centro=True)
        _texto(cells[1], str(it.get("numero", "")), size=9, cor=MUT)
        cells[2].paragraphs[0].paragraph_format.left_indent = Cm((nivel - 1) * 0.4)
        r = _texto(cells[2], it.get("texto", ""), bold=nivel == 1, size=10 if nivel == 1 else 9.5,
                   cor=SLATE if nivel == 1 else None)
        for c in (3, 4, 5):
            _texto(cells[c], "___/___", size=9, cor=MUT, centro=True)
    for col, w in zip(t.columns, LARG):   # grade da tabela (LibreOffice/Google Docs leem daqui)
        col.width = w
    for row in t.rows:
        for cell, w in zip(row.cells, LARG):
            cell.width = w
    return t


def _page_number(paragraph):
    run = paragraph.add_run()
    for t, val in (("begin", None), ("instr", "PAGE"), ("end", None)):
        if t == "instr":
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = " PAGE "
        else:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), t)
        run._r.append(el)


def build(data, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for m in ("top", "bottom", "left", "right"):
        setattr(sec, f"{m}_margin", Cm(2))

    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(11)
    h1 = doc.styles["Heading 1"]; h1.font.name = "Calibri"; h1.font.size = Pt(15); h1.font.color.rgb = AZUL

    # Cabeçalho
    if LOGO:
        try:
            doc.add_picture(LOGO, width=Cm(1.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            pass
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Edital Verticalizado"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = SLATE
    sub = " · ".join(x for x in [data.get("concurso"), data.get("cargo")] if x)
    if sub:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(sub); rr.font.size = Pt(12); rr.font.color.rgb = AZUL
    meta = " · ".join(x for x in [data.get("orgao"), data.get("data_edital")] if x)
    if meta:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(meta); rr.font.size = Pt(9.5); rr.font.color.rgb = MUT

    # Quadro resumo
    doc.add_paragraph()
    qh = doc.add_paragraph(); rq = qh.add_run("Quadro-resumo das disciplinas"); rq.bold = True
    rq.font.size = Pt(13); rq.font.color.rgb = SLATE
    grupos = []   # [(nome do bloco ou "", [disciplinas])] na ordem do edital
    for d in data.get("disciplinas", []):
        b = BLOCOS.get(d.get("bloco", ""), "")
        if not grupos or grupos[-1][0] != b:
            grupos.append((b, []))
        grupos[-1][1].append(d)
    for b, discs in grupos:
        if b:
            pb = doc.add_paragraph(); rb = pb.add_run(f"{b} ({sum(len(d.get('itens', [])) for d in discs)} itens)")
            rb.bold = True; rb.font.color.rgb = AZUL
        for d in discs:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{d.get('nome','')} ").bold = True
            p.add_run(f"({len(d.get('itens', []))} itens)").font.color.rgb = MUT
    total = sum(len(d.get("itens", [])) for d in data.get("disciplinas", []))
    tot = doc.add_paragraph(); rt = tot.add_run(f"Total: {total} itens para acompanhar.")
    rt.italic = True; rt.font.color.rgb = MUT
    dica = doc.add_paragraph()
    rd = dica.add_run("Como anotar: marque ☐ ao concluir o item e, em cada fase, escreva certas/resolvidas "
                      "(ex.: 7/8). Meta: 70% de acerto ou mais.")
    rd.font.size = Pt(9.5); rd.font.color.rgb = MUT

    # Disciplinas (uma por página), agrupadas por bloco
    for b, discs in grupos:
        for i, d in enumerate(discs):
            doc.add_page_break()
            if b and i == 0:
                pb = doc.add_paragraph(); rb = pb.add_run(b.upper())
                rb.bold = True; rb.font.size = Pt(10); rb.font.color.rgb = MUT
            doc.add_heading(d.get("nome", ""), level=1)
            _tabela_itens(doc, d.get("itens", []))

    # Rodapé com número de página
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Edital Verticalizado · Simião Cavalcante   |   ")
    fr.font.size = Pt(8); fr.font.color.rgb = MUT
    _page_number(fp)

    doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("json"); ap.add_argument("--out", "-o", required=True)
    ap.add_argument("--sem-logo", action="store_true", help="não inserir logo no topo")
    a = ap.parse_args()
    if a.sem_logo:
        global LOGO
        LOGO = ""
    data = json.loads(Path(a.json).read_text(encoding="utf-8"))
    build(data, a.out)
    n = sum(len(d["itens"]) for d in data.get("disciplinas", []))
    print(f"OK -> {a.out} | {len(data.get('disciplinas', []))} disciplinas, {n} itens")


if __name__ == "__main__":
    main()
